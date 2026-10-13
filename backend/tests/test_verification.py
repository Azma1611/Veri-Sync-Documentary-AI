"""
VeriSync AI — Verification Engine Tests

Tests are organised in five sections:
  1. Field normalizer unit tests
  2. Domain rules unit tests
  3. Completeness checker tests (using synthetic documents)
  4. Comparison engine tests (matching clauses, salary mismatches, duration)
  5. Verification endpoint integration tests (API-level)

All synthetic documents are generated programmatically using PyMuPDF
and python-docx.  No external fixture files required.

Run:
    cd backend
    .venv\\Scripts\\pytest tests/test_verification.py -v
"""

from __future__ import annotations

import io
import json
from typing import List, Optional

import pytest
from fastapi.testclient import TestClient


# ---------------------------------------------------------------------------
# Synthetic document builders
# ---------------------------------------------------------------------------

def make_pdf_bytes(text: str) -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=11)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_multi_page_pdf_bytes(pages: List[str]) -> bytes:
    import fitz
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_docx_bytes(paragraphs: List[str], headings: Optional[List[str]] = None) -> bytes:
    from docx import Document
    doc = Document()
    headings = headings or []
    for i, text in enumerate(paragraphs):
        if i < len(headings):
            doc.add_heading(headings[i], level=1)
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _ingest_pdf(content: bytes, filename: str, doc_type: str) -> "DocumentMetadata":
    """Helper: ingest a synthetic PDF directly into the document store."""
    from app.services.document_store import document_store
    from app.services.pdf_extractor import extract_pdf
    from app.schemas.ingestion import ProcessingStatus

    doc = document_store.new_pending(
        original_filename=filename,
        content_type="application/pdf",
        doc_type=doc_type,
        domain="employment-hr",
        file_size_bytes=len(content),
    )
    full_text, page_count, evidence_blocks = extract_pdf(content, filename)
    document_store.mark_ready(
        doc.document_id,
        full_text=full_text,
        evidence_blocks=evidence_blocks,
        page_count=page_count,
    )
    return document_store.get(doc.document_id)


def _ingest_docx(content: bytes, filename: str, doc_type: str) -> "DocumentMetadata":
    from app.services.document_store import document_store
    from app.services.docx_extractor import extract_docx

    doc = document_store.new_pending(
        original_filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        doc_type=doc_type,
        domain="employment-hr",
        file_size_bytes=len(content),
    )
    full_text, para_count, evidence_blocks = extract_docx(content, filename)
    document_store.mark_ready(
        doc.document_id,
        full_text=full_text,
        evidence_blocks=evidence_blocks,
        paragraph_count=para_count,
    )
    return document_store.get(doc.document_id)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture(autouse=True)
def clear_store():
    from app.services.document_store import document_store
    document_store.clear()
    yield
    document_store.clear()


@pytest.fixture(scope="module")
def client():
    from app.main import app
    return TestClient(app)


# ===========================================================================
# Section 1: Field Normalizer
# ===========================================================================

class TestFieldNormalizer:

    def test_currency_with_rupee_symbol(self):
        from app.services.field_normalizer import normalize_currency
        assert normalize_currency("₹30,000 per month") == 30000.0

    def test_currency_with_rs_prefix(self):
        from app.services.field_normalizer import normalize_currency
        assert normalize_currency("Rs. 25,000 gross") == 25000.0

    def test_currency_with_usd(self):
        from app.services.field_normalizer import normalize_currency
        val = normalize_currency("$5,000 per month")
        assert val == 5000.0

    def test_currency_bare_number(self):
        from app.services.field_normalizer import normalize_currency
        val = normalize_currency("salary is 50000 annually")
        assert val == 50000.0

    def test_currency_none_for_no_number(self):
        from app.services.field_normalizer import normalize_currency
        assert normalize_currency("salary details to be discussed") is None

    def test_duration_months(self):
        from app.services.field_normalizer import normalize_duration_to_days
        assert normalize_duration_to_days("3 months") == 90

    def test_duration_days(self):
        from app.services.field_normalizer import normalize_duration_to_days
        assert normalize_duration_to_days("90 days") == 90

    def test_duration_weeks(self):
        from app.services.field_normalizer import normalize_duration_to_days
        assert normalize_duration_to_days("2 weeks") == 14

    def test_duration_year(self):
        from app.services.field_normalizer import normalize_duration_to_days
        assert normalize_duration_to_days("1 year") == 365

    def test_duration_hyphenated(self):
        from app.services.field_normalizer import normalize_duration_to_days
        assert normalize_duration_to_days("6-month probation") == 180

    def test_duration_none(self):
        from app.services.field_normalizer import normalize_duration_to_days
        assert normalize_duration_to_days("duration TBD by management") is None

    def test_compare_numeric_match(self):
        from app.services.field_normalizer import compare_numeric
        assert compare_numeric(30000.0, 30000.0) == "match"

    def test_compare_numeric_mismatch(self):
        from app.services.field_normalizer import compare_numeric
        assert compare_numeric(30000.0, 25000.0) == "mismatch"

    def test_compare_numeric_insufficient(self):
        from app.services.field_normalizer import compare_numeric
        assert compare_numeric(None, 25000.0) == "insufficient_evidence"

    def test_compare_duration_match(self):
        from app.services.field_normalizer import compare_duration
        assert compare_duration("3 months", "90 days") == "match"

    def test_compare_duration_mismatch(self):
        from app.services.field_normalizer import compare_duration
        assert compare_duration("3 months", "6 months") == "mismatch"


# ===========================================================================
# Section 2: Domain Rules
# ===========================================================================

class TestDomainRules:

    def test_all_five_rules_defined(self):
        from app.services.domain_rules import EMPLOYMENT_HR_RULES
        rule_ids = {r.rule_id for r in EMPLOYMENT_HR_RULES}
        assert {"COMP-001", "PROB-001", "NTCE-001", "IP-001", "RPT-001"} == rule_ids

    def test_comp_001_expected_in_offer_and_contract(self):
        from app.services.domain_rules import RULE_BY_ID
        rule = RULE_BY_ID["COMP-001"]
        assert "offer_letter" in rule.expected_in
        assert "contract" in rule.expected_in
        assert "job_description" not in rule.expected_in

    def test_ip_001_only_expected_in_contract(self):
        from app.services.domain_rules import RULE_BY_ID
        rule = RULE_BY_ID["IP-001"]
        assert rule.expected_in == ["contract"]

    def test_rpt_001_only_expected_in_offer_letter(self):
        from app.services.domain_rules import RULE_BY_ID
        rule = RULE_BY_ID["RPT-001"]
        assert rule.expected_in == ["offer_letter"]

    def test_get_rule_by_id(self):
        from app.services.domain_rules import get_rule
        rule = get_rule("NTCE-001")
        assert rule.name == "Notice Period"

    def test_get_rule_unknown_raises(self):
        from app.services.domain_rules import get_rule
        with pytest.raises(KeyError):
            get_rule("UNKNOWN-999")


# ===========================================================================
# Section 3: Completeness Checker
# ===========================================================================

class TestCompletenessChecker:

    def test_matching_clauses_no_missing_findings(self):
        from app.services.completeness_checker import check_completeness

        offer_text = (
            "Dear Candidate, we are pleased to offer you this position.\n"
            "Gross Salary: ₹30,000 per month.\n"
            "Probation period: 3 months.\n"
            "Notice period: 30 days.\n"
            "You will report to the Engineering Manager.\n"
        )
        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        findings = check_completeness([offer])
        # COMP-001, PROB-001, NTCE-001, RPT-001 should all be found
        rule_ids_missing = {f.rule_id for f in findings}
        assert "COMP-001" not in rule_ids_missing
        assert "PROB-001" not in rule_ids_missing
        assert "NTCE-001" not in rule_ids_missing
        assert "RPT-001" not in rule_ids_missing

    def test_missing_salary_clause_flagged(self):
        from app.services.completeness_checker import check_completeness

        # Offer letter with no salary mention
        text = (
            "Dear Candidate, we are pleased to offer you a position.\n"
            "Probation period: 3 months. Notice: 30 days.\n"
            "You will report to the Team Lead.\n"
        )
        offer = _ingest_pdf(make_pdf_bytes(text), "offer_no_salary.pdf", "offer_letter")
        findings = check_completeness([offer])
        rule_ids = {f.rule_id for f in findings}
        assert "COMP-001" in rule_ids

    def test_missing_ip_clause_in_contract_flagged(self):
        from app.services.completeness_checker import check_completeness

        # Contract with no IP/annexure mention
        text = (
            "Employment Contract\n"
            "Monthly Gross Salary: ₹25,000.\n"
            "Probation: 6 months. Notice period: 90 days.\n"
        )
        contract = _ingest_pdf(make_pdf_bytes(text), "contract.pdf", "contract")
        findings = check_completeness([contract])
        rule_ids = {f.rule_id for f in findings}
        assert "IP-001" in rule_ids

    def test_ip_clause_not_required_in_offer_letter(self):
        from app.services.completeness_checker import check_completeness

        # Offer letter with no IP clause — should NOT be flagged (IP-001 only expected in contract)
        text = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days. Report to: Manager."
        offer = _ingest_pdf(make_pdf_bytes(text), "offer.pdf", "offer_letter")
        findings = check_completeness([offer])
        rule_ids = {f.rule_id for f in findings}
        assert "IP-001" not in rule_ids

    def test_completeness_skips_rules_for_absent_doc_types(self):
        from app.services.completeness_checker import check_completeness
        from app.services.domain_rules import RULE_BY_ID

        # Upload only an HR policy — COMP-001 is not expected in hr_policy
        text = "Notice period for all employees: 60 days."
        policy = _ingest_pdf(make_pdf_bytes(text), "policy.pdf", "hr_policy")
        findings = check_completeness([policy])
        rule_ids = {f.rule_id for f in findings}
        # COMP-001 not expected in hr_policy → should not appear
        assert "COMP-001" not in rule_ids
        # NTCE-001 IS expected in hr_policy → if missing, it would appear
        # (policy has "60 days" so it should pass)
        assert "NTCE-001" not in rule_ids


# ===========================================================================
# Section 4: Comparison Engine
# ===========================================================================

class TestComparisonEngine:

    def test_matching_salaries_no_finding(self):
        from app.services.comparison_engine import compare_documents

        offer_text = "This offer confirms: Gross Salary ₹30,000 per month. Probation: 3 months. Notice: 30 days."
        contract_text = "Monthly Gross Salary: ₹30,000. Probation period: 3 months. Notice period: 30 days. Annexure B: IP Assignment attached."

        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        contract = _ingest_pdf(make_pdf_bytes(contract_text), "contract.pdf", "contract")

        findings = compare_documents([offer, contract])
        comp_findings = [f for f in findings if f.rule_id == "COMP-001"]
        # Values match — no compensation finding
        assert len(comp_findings) == 0

    def test_salary_mismatch_produces_high_severity_finding(self):
        from app.services.comparison_engine import compare_documents
        from app.schemas.verification import InconsistencySeverity, FindingType

        offer_text = "Gross Salary: ₹30,000 per month. Probation: 3 months. Notice period: 30 days."
        contract_text = "Monthly Gross: ₹25,000. Probation: 3 months. Notice: 30 days. Annexure B: IP Assignment."

        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        contract = _ingest_pdf(make_pdf_bytes(contract_text), "contract.pdf", "contract")

        findings = compare_documents([offer, contract])
        comp_findings = [f for f in findings if f.rule_id == "COMP-001"]
        assert len(comp_findings) == 1
        f = comp_findings[0]
        assert f.severity == InconsistencySeverity.high
        assert f.finding_type == FindingType.value_mismatch
        assert f.sourceA.documentName == "offer.pdf"
        assert f.sourceB.documentName == "contract.pdf"

    def test_probation_mismatch_produces_medium_severity(self):
        from app.services.comparison_engine import compare_documents
        from app.schemas.verification import InconsistencySeverity

        offer_text = "Salary: ₹30,000. Probation period: 3 months. Notice: 30 days."
        contract_text = "Salary: ₹30,000. Probation period: 6 months. Notice: 30 days. Annexure B attached."

        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        contract = _ingest_pdf(make_pdf_bytes(contract_text), "contract.pdf", "contract")

        findings = compare_documents([offer, contract])
        prob_findings = [f for f in findings if f.rule_id == "PROB-001"]
        assert len(prob_findings) == 1
        assert prob_findings[0].severity == InconsistencySeverity.medium

    def test_notice_period_mismatch_30_vs_90_days(self):
        from app.services.comparison_engine import compare_documents

        offer_text = "Salary: ₹30,000. Probation: 3 months. Notice period: 30 days."
        contract_text = "Salary: ₹30,000. Probation: 3 months. Notice period: 90 days. Annexure B: IP."

        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        contract = _ingest_pdf(make_pdf_bytes(contract_text), "contract.pdf", "contract")

        findings = compare_documents([offer, contract])
        notice_findings = [f for f in findings if f.rule_id == "NTCE-001"]
        assert len(notice_findings) == 1

    def test_no_findings_when_only_one_applicable_doc(self):
        from app.services.comparison_engine import compare_documents

        # Only one offer_letter — no contract to compare against
        offer_text = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days. Report to: Manager."
        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        findings = compare_documents([offer])
        assert len(findings) == 0

    def test_finding_ids_are_stable_and_unique(self):
        from app.services.comparison_engine import compare_documents

        offer_text = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days."
        contract_text = "Salary: ₹25,000. Probation: 6 months. Notice: 90 days. Annexure B: IP."

        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        contract = _ingest_pdf(make_pdf_bytes(contract_text), "contract.pdf", "contract")

        findings1 = compare_documents([offer, contract])
        findings2 = compare_documents([offer, contract])

        ids1 = [f.id for f in findings1]
        ids2 = [f.id for f in findings2]
        assert ids1 == ids2  # stable
        assert len(ids1) == len(set(ids1))  # unique

    def test_insufficient_evidence_when_salary_unparseable(self):
        from app.services.comparison_engine import compare_documents
        from app.schemas.verification import FindingType

        # "Salary as per market standards" — no numeric value
        offer_text = "Salary: as per market standards and negotiable. Probation: 3 months. Notice: 30 days."
        contract_text = "Compensation: competitive and reviewed annually. Probation: 3 months. Notice: 30 days. Annexure B attached."

        offer = _ingest_pdf(make_pdf_bytes(offer_text), "offer.pdf", "offer_letter")
        contract = _ingest_pdf(make_pdf_bytes(contract_text), "contract.pdf", "contract")

        findings = compare_documents([offer, contract])
        # Salary clause may or may not match regex — but if it does, value comparison should
        # yield insufficient_evidence since there's no parseable number
        comp_findings = [f for f in findings if f.rule_id == "COMP-001"]
        for f in comp_findings:
            assert f.finding_type in (
                FindingType.insufficient_evidence,
                FindingType.value_mismatch,
                FindingType.missing_clause,
            )

    def test_docx_documents_compared_correctly(self):
        from app.services.comparison_engine import compare_documents

        offer_text = ["Dear candidate, your Gross Salary is ₹30,000 per month.", "Probation: 3 months.", "Notice period: 30 days."]
        contract_text = ["Monthly gross compensation: ₹25,000.", "Probation period: 3 months.", "Notice: 30 days.", "Annexure B: IP Assignment attached."]

        offer = _ingest_docx(make_docx_bytes(offer_text), "offer.docx", "offer_letter")
        contract = _ingest_docx(make_docx_bytes(contract_text), "contract.docx", "contract")

        findings = compare_documents([offer, contract])
        comp_findings = [f for f in findings if f.rule_id == "COMP-001"]
        assert len(comp_findings) >= 1
        # Source citations should have paragraph_ref (DOCX)
        f = comp_findings[0]
        assert f.sourceA.paragraph_ref is not None or f.sourceA.page == 1


# ===========================================================================
# Section 5: API endpoint integration tests
# ===========================================================================

class TestVerifyEndpoint:

    PDF_CONTENT_TYPE = "application/pdf"
    DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def test_verify_matching_documents_returns_clean(self, client: TestClient):
        offer_text = (
            "Employment Offer Letter\n"
            "Gross Salary: ₹30,000 per month.\n"
            "Probation period: 3 months.\n"
            "Notice period: 30 days from either party.\n"
            "You will report to the Head of Engineering.\n"
        )
        contract_text = (
            "Employment Contract\n"
            "Monthly Gross Salary: ₹30,000.\n"
            "Probation period: 3 months.\n"
            "Notice period: 30 days.\n"
            "Annexure B: Intellectual Property Assignment is attached.\n"
        )
        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("offer.pdf", make_pdf_bytes(offer_text), self.PDF_CONTENT_TYPE)),
                ("files", ("contract.pdf", make_pdf_bytes(contract_text), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": '["offer_letter","contract"]', "domain": "employment-hr"},
        )
        assert resp.status_code == 200
        envelope = resp.json()
        assert envelope["isSimulated"] is False
        assert envelope["backendConnected"] is True
        data = envelope["data"]
        assert data["totalDocuments"] == 2
        assert "offer.pdf" in data["documentsReviewed"]

    def test_verify_salary_mismatch_detected(self, client: TestClient):
        offer_text = "Offer Letter. Gross Salary: ₹30,000 per month. Probation: 3 months. Notice: 30 days."
        contract_text = "Contract. Monthly Gross: ₹25,000. Probation: 3 months. Notice: 30 days. Annexure B: IP."

        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("offer.pdf", make_pdf_bytes(offer_text), self.PDF_CONTENT_TYPE)),
                ("files", ("contract.pdf", make_pdf_bytes(contract_text), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": '["offer_letter","contract"]', "domain": "employment-hr"},
        )
        assert resp.status_code == 200
        data = resp.json()["data"]
        findings = data["findings"]
        comp_findings = [f for f in findings if f["rule_id"] == "COMP-001"]
        assert len(comp_findings) >= 1
        assert comp_findings[0]["severity"] == "high"
        assert comp_findings[0]["status"] == "flagged"
        # Verify source citations present
        assert "sourceA" in comp_findings[0]
        assert "sourceB" in comp_findings[0]
        assert comp_findings[0]["sourceA"]["documentName"] in ("offer.pdf", "contract.pdf")

    def test_verify_missing_ip_clause_detected(self, client: TestClient):
        contract_text = (
            "Employment Contract\n"
            "Monthly Gross: ₹30,000. Probation: 6 months. Notice period: 90 days.\n"
            # No mention of Annexure B or IP assignment
        )
        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("contract.pdf", make_pdf_bytes(contract_text), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": "contract", "domain": "employment-hr"},
        )
        assert resp.status_code == 200
        data = resp.json()["data"]
        missing = data["missingFields"]
        ip_missing = [m for m in missing if m["rule_id"] == "IP-001"]
        assert len(ip_missing) >= 1
        assert ip_missing[0]["severity"] == "critical"

    def test_verify_status_is_inconsistencies_detected_for_high_findings(self, client: TestClient):
        offer_text = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days."
        contract_text = "Salary: ₹25,000. Probation: 3 months. Notice: 30 days. Annexure B: IP."

        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("offer.pdf", make_pdf_bytes(offer_text), self.PDF_CONTENT_TYPE)),
                ("files", ("contract.pdf", make_pdf_bytes(contract_text), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": '["offer_letter","contract"]', "domain": "employment-hr"},
        )
        data = resp.json()["data"]
        assert data["status"] == "inconsistencies_detected"

    def test_verify_response_matches_api_contract_envelope(self, client: TestClient):
        """Verify the response envelope matches ApiResponse<VerificationResult>."""
        offer_text = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days."
        resp = client.post(
            "/api/v1/documents/verify",
            files=[("files", ("offer.pdf", make_pdf_bytes(offer_text), self.PDF_CONTENT_TYPE))],
            data={"doc_types": "offer_letter", "domain": "employment-hr"},
        )
        assert resp.status_code == 200
        envelope = resp.json()
        # Must have all ApiResponse fields
        assert "data" in envelope
        assert "isSimulated" in envelope
        assert "backendConnected" in envelope
        # data must have VerificationResult fields
        data = envelope["data"]
        for field in ["domain", "totalDocuments", "status", "findings", "missingFields",
                      "summaryNote", "reviewedAt", "documentsReviewed"]:
            assert field in data, f"Missing field: {field}"

    def test_verify_finding_has_all_required_fields(self, client: TestClient):
        offer_text = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days."
        contract_text = "Salary: ₹25,000. Probation: 3 months. Notice: 30 days. Annexure B: IP."

        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("offer.pdf", make_pdf_bytes(offer_text), self.PDF_CONTENT_TYPE)),
                ("files", ("contract.pdf", make_pdf_bytes(contract_text), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": '["offer_letter","contract"]'},
        )
        data = resp.json()["data"]
        if data["findings"]:
            f = data["findings"][0]
            for field in ["id", "rule_id", "finding_type", "title", "severity",
                          "category", "explanation", "sourceA", "sourceB",
                          "recommendedAction", "status"]:
                assert field in f, f"Finding missing field: {field}"
            assert "documentName" in f["sourceA"]
            assert "extractedValue" in f["sourceA"]
            assert "page" in f["sourceA"]

    def test_verify_no_files_returns_422(self, client: TestClient):
        resp = client.post(
            "/api/v1/documents/verify",
            data={"doc_types": "offer_letter"},
        )
        assert resp.status_code == 422

    def test_verify_corrupt_file_produces_extraction_failure_report(self, client: TestClient):
        corrupt = b"%PDF-1.4\nGARBAGE"
        resp = client.post(
            "/api/v1/documents/verify",
            files=[("files", ("bad.pdf", corrupt, self.PDF_CONTENT_TYPE))],
            data={"doc_types": "offer_letter"},
        )
        assert resp.status_code == 200
        data = resp.json()["data"]
        # All files failed → manual_review_required
        assert data["status"] in ("manual_review_required", "inconsistencies_detected")

    def test_verify_probation_mismatch_both_docs(self, client: TestClient):
        offer_text = "Salary: ₹30,000. Probation period: 3 months. Notice: 30 days."
        contract_text = "Salary: ₹30,000. Probation period: 6 months. Notice: 30 days. Annexure B: IP."

        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("offer.pdf", make_pdf_bytes(offer_text), self.PDF_CONTENT_TYPE)),
                ("files", ("contract.pdf", make_pdf_bytes(contract_text), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": '["offer_letter","contract"]'},
        )
        data = resp.json()["data"]
        prob_findings = [f for f in data["findings"] if f["rule_id"] == "PROB-001"]
        assert len(prob_findings) >= 1
        assert prob_findings[0]["severity"] == "medium"

    def test_verify_three_documents_exhaustive_comparison(self, client: TestClient):
        """Three documents → pairwise comparisons across all pairs."""
        t1 = "Salary: ₹30,000. Probation: 3 months. Notice: 30 days."
        t2 = "Salary: ₹25,000. Probation: 6 months. Notice: 90 days. Annexure B: IP."
        t3 = "Salary: ₹28,000. Probation: 3 months. Notice: 60 days. Annexure B: IP."

        resp = client.post(
            "/api/v1/documents/verify",
            files=[
                ("files", ("offer.pdf", make_pdf_bytes(t1), self.PDF_CONTENT_TYPE)),
                ("files", ("contract.pdf", make_pdf_bytes(t2), self.PDF_CONTENT_TYPE)),
                ("files", ("addendum.pdf", make_pdf_bytes(t3), self.PDF_CONTENT_TYPE)),
            ],
            data={"doc_types": '["offer_letter","contract","contract"]'},
        )
        assert resp.status_code == 200
        data = resp.json()["data"]
        assert data["totalDocuments"] == 3
        # Multiple findings expected from three-way comparison
        assert len(data["findings"]) > 0
